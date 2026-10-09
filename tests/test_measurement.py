"""The measurement harness, run on labels whose answer is known by construction.

Gold does not exist yet; the harness must already be right when it arrives.
Nothing here touches the network or the 30 MB crawl: the stance model is a fake
and the company library is a handful of pages.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from quantum_gw.providers.base import LLMResponse
from quantum_gw.settings import load_settings
from quantum_gw.verification import llm_judge

ROOT = Path(__file__).resolve().parents[1]


def _load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "tools" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def ev():
    sys.path.insert(0, str(ROOT / "tools"))
    return _load("evaluate_gold")


@pytest.fixture(scope="module")
def ma(ev):
    return _load("measure_all")


CLAIM = "Năm 2024, tổng lượng nước tiêu thụ tại hai nhà máy giảm 8% so với năm 2022."
PASSAGES = {
    "d1:u5:aa": "Năm 2024, lượng nước tiêu thụ tại hai nhà máy là 1,21 triệu m3, giảm 8% so với năm 2022.",
    "d1:u9:bb": "Công ty tổ chức ngày hội trồng cây cho người lao động tại văn phòng Hà Nội.",
    "d1:u12:cc": "Chương trình tiết kiệm nước tiếp tục được duy trì tại các nhà máy trong năm.",
}


def _row(labels_by_chunk, verdict="PARTIALLY_SUPPORTED", company="x", year=2024):
    candidates = [
        {"rank": i, "evidence_id": f"{cid}:evidence:{i}", "retrieval_score": 0.5, "text": text,
         "document_type": "sustainability_report", "publication_year": year, "unit_index": i}
        for i, (cid, text) in enumerate(PASSAGES.items(), start=1)
    ]
    return {
        "claim_id": "d0:u1:zz:claim:1", "company_id": company, "ticker": "X", "document_id": "d0",
        "document_type": "sustainability_report", "publication_year": year, "unit_index": 1,
        "split": "test", "text": CLAIM, "evidence_candidates": candidates,
        "labels": {"is_claim": "yes", "evidence": {
            f"{cid}:evidence:{i}": labels_by_chunk[cid]
            for i, cid in enumerate(PASSAGES, start=1) if cid in labels_by_chunk},
            "verdict": verdict, "labeler": "test"},
    }


LABELS = {"d1:u5:aa": "supports", "d1:u9:bb": "not_relevant", "d1:u12:cc": "partial"}


def _settings(stance="off", tmp_path=None):
    s = load_settings("configs/default.yaml")
    s.verification.llm_stance = stance
    if tmp_path is not None:
        s.verification.llm_cache_dir = str(tmp_path / "cache")
    return s


# --- stance, per pair --------------------------------------------------------

def test_every_labelled_pair_is_scored_and_not_relevant_counts_as_context(ev):
    result = ev.stance_metrics([_row(LABELS)], _settings())
    assert result["rows"] == 3
    humans = sorted(r["human"] for r in result["records"])
    assert humans == ["CONTEXT", "PARTIAL", "SUPPORTS"]
    assert 0.0 <= result["accuracy"]["mean"] <= 1.0
    assert result["model_vs_rules_on_escalated"] is None     # no model when off


def test_the_model_is_compared_with_the_rules_on_the_very_same_pairs(ev, tmp_path, monkeypatch):
    class Gateway:
        def route(self, task_type=None):
            return ["fake:m"]

        def generate(self, messages, response_schema=None, task_type=None):
            return LLMResponse(text="{}", parsed={"relation": "CONTEXT", "reason": "r"},
                               provider="fake", model="m")

    monkeypatch.setattr(llm_judge, "ModelGateway", Gateway)
    result = ev.stance_metrics([_row(LABELS)], _settings("on_ambiguous", tmp_path))
    escalated = result["model_vs_rules_on_escalated"]
    assert escalated and escalated["n"] == sum(1 for r in result["records"] if r["method"] == "llm")
    for record in result["records"]:
        if record["method"] != "llm":
            assert record["rules"] == record["system"]


def test_unlabelled_or_non_claim_rows_are_left_out(ev):
    row = _row(LABELS)
    row["labels"]["is_claim"] = "no"
    assert ev.stance_metrics([row], _settings())["rows"] == 0


# --- retrieval over the library ---------------------------------------------

def _library(pages_by_company):
    def fake(companies):
        return {c: pages_by_company.get(c, []) for c in companies}
    return fake


def _page(cid, text, year=2024):
    return {"chunk_id": cid, "document_id": cid.split(":")[0], "ticker": "X",
            "document_type": "sustainability_report", "publication_year": year,
            "unit_index": 1, "text": text}


def test_the_labelled_page_is_found_in_the_library_and_later_pages_are_invisible(ev, monkeypatch):
    pages = [_page(cid, text) for cid, text in PASSAGES.items()]
    pages.append(_page("d9:u1:later", PASSAGES["d1:u5:aa"], year=2026))  # same words, published later
    monkeypatch.setattr(ev, "_library", _library({"x": pages}))
    result = ev.retrieval_over_library([_row(LABELS)], _settings())
    assert result["rows"] == 1
    assert result["library_pages"] == 4
    assert result["recall_at"]["@30"]["mean"] == 1.0          # both relevant pages are reachable
    assert result["mrr"]["mean"] == 1.0                        # the confirming page ranks first
    assert "verdict_accuracy" in result


# --- decisions follow the registered thresholds ------------------------------

def _rules():
    import yaml
    return yaml.safe_load((ROOT / "configs" / "measurement_decisions.yaml").read_text(encoding="utf-8"))


def test_a_model_below_the_registered_precision_keeps_its_guard(ma):
    results = {"stance": {"model": {"by_method": {"llm": {"precision": {
        "SUPPORTS": {"mean": 0.7, "n": 40, "ci95": [0.55, 0.83]}}}},
        "model_vs_rules_on_escalated": {"n": 40, "difference": {"mean": 0.1, "n": 40, "ci95": [0.02, 0.2]}}}}}
    made = ma.decide(results, _rules())
    decisions = {d["rule"]: d for d in made}
    assert decisions["lift_guard_supports"]["decision"] == "GIỮ CHỐT CHẶN"
    assert decisions["stance_model_keep_on"]["decision"] == "GIỮ BẬT"
    assert decisions["lift_guard_contradicts"]["decision"] == "GIỮ CHỐT CHẶN"
    assert {d["decision"] for d in made if d["rule"] == "dense_default"} == {"KHÔNG ĐO ĐƯỢC"}


def test_dense_retrieval_is_switched_on_only_past_both_registered_bars(ma):
    def run(recall, seconds):
        return {"rows": 2, "seconds_per_claim": seconds,
                "per_row": [{"claim_id": "a", "recall@5": recall}, {"claim_id": "b", "recall@5": recall}]}
    fast_gain = {"retrieval": {"lite": run(0.4, 1.0), "vn-embedding": run(0.6, 1.5)}}
    slow_gain = {"retrieval": {"lite": run(0.4, 1.0), "vn-embedding": run(0.6, 3.0)}}
    small_gain = {"retrieval": {"lite": run(0.4, 1.0), "vn-embedding": run(0.45, 1.1)}}

    def decide(results):
        return next(d["decision"] for d in ma.decide(results, _rules())
                    if d["rule"] == "dense_default" and d["value"].startswith("vn-embedding"))
    assert decide(fast_gain) == "BẬT MẶC ĐỊNH"
    assert decide(slow_gain) == "CHƯA BẬT"
    assert decide(small_gain) == "CHƯA BẬT"


# --- returned workbooks --------------------------------------------------------

def test_the_hpg_queue_and_the_timing_study_come_back_as_numbers(tmp_path, monkeypatch):
    sys.path.insert(0, str(ROOT / "tools"))
    lp = _load("labeling_pack")
    packs = _load("make_handoff_packs")
    monkeypatch.setattr(lp, "BENCHMARK", tmp_path)

    queue = tmp_path / "queue.xlsx"
    packs.hpg_queue(queue)
    wb = load_workbook(queue)
    ws = wb["Hàng đợi HPG"]
    for r, answer in zip(range(4, 18), ["Có"] * 9 + ["Không"] * 5, strict=True):
        ws.cell(r, 9, answer)
    ws.cell(13, 10, "giải thích quy trình / kiến thức chung")
    wb.save(queue)
    import json
    report = json.loads(lp.import_hpg_queue(queue).read_text(encoding="utf-8"))
    assert report["top14"] == {"answered": 14, "yes": 9, "precision": round(9 / 14, 4)}

    timing = tmp_path / "timing.xlsx"
    packs.timing_study(timing)
    wb = load_workbook(timing)
    ws = wb["Bài đo"]
    ws.cell(4, 6, "14:05")
    ws.cell(4, 7, "14:17")
    wb.save(timing)
    report = json.loads(lp.import_timing(timing).read_text(encoding="utf-8"))
    assert report["completed"] == 1
    assert next(i for i in report["items"] if i["n"] == 1)["minutes"] == 12.0
