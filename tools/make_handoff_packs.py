#!/usr/bin/env python3
"""Build the two work packs of 28/09/2026 — one for Quỳnh, one for Thảo.

Everything either of them needs is inside their own folder, so neither has to
open the repository or a terminal:

    Goi_Quynh/  gold workbooks (lô 1–3), the HPG queue check, the timing study,
                the two HPG reports the queue and the study refer to
    Goi_Thao/   her copy of lô 1 (the κ batch), the three-rule form (Word),
                the official pages each rule rests on, candidate examples
                from real reports, the wording review, the slides

    python tools/make_handoff_packs.py build      # folders
    python tools/make_handoff_packs.py zip        # after the guide PDF is compiled

Legal page extracts are cut from the signed originals in data/legal/raw, not
from OCR, so Thảo reads the official text. Everything derived from OCR (the
QĐ 13 facility lines, clause drafts in the form) is labelled as needing a check.
"""

from __future__ import annotations

import argparse
import json
import random
import re
import shutil
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import labeling_pack as lp  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "06-operations" / "handoff_2026-09-28"
QUYNH = OUT / "Goi_Quynh"
THAO = OUT / "Goi_Thao"
GUIDE = OUT / "HuongDan_Thao_Quynh_2026-09-28.pdf"

SESSIONS = ROOT / "data" / "gold" / "sessions"
LOTS = [("2026-09-28_lo1_kappa.jsonl", "01_Gold_Lo1_kappa_20cau"),
        ("2026-09-28_lo2.jsonl", "02_Gold_Lo2_30cau"),
        ("2026-09-28_lo3.jsonl", "03_Gold_Lo3_30cau")]
HPG_RUN = ROOT / ".quantum" / "runs" / "185480a0c8b747c4" / "result.json"
HPG_PDF = ROOT / "data" / "real_cases" / "sources" / "originals"
LEGAL = ROOT / "data" / "legal" / "raw"
CRAWL = ROOT / "data" / "crawl" / "vn30" / "normalized"
SLIDES = ROOT / "docs" / "07-presentation" / "GREENSCAN.zip"
STAMP_FONT = "C:/Windows/Fonts/arial.ttf"  # the built-in PDF fonts have no Vietnamese diacritics

GREEN = "1B7F5A"
HEAD = PatternFill("solid", fgColor=GREEN)
TODO = PatternFill("solid", fgColor="FFF4CC")
WRAP = Alignment(wrap_text=True, vertical="top")

# Pages of the signed originals each rule rests on (1-based, found by mapping
# OCR hits back to pages; see the guide, section Thảo).
LEGAL_PAGES = {
    "QD13-2024_Dieu1-3_PhuLucI_dau-PhuLucII": ("QD13-2024-QD-TTg", [1, 2, 3, 4, 5]),
    "QD13-2024_trang_co_so_DN_niem_yet": ("QD13-2024-QD-TTg",
                                          [37, 45, 46, 47, 59, 61, 62, 66, 68, 69, 93, 156, 169, 175]),
    "ND06-2022_Dieu11_kiem_ke": ("ND06-2022-ND-CP", [1, 9, 10, 11]),
    "ND119-2025_sua_Dieu11_va_hieu_luc": ("ND119-2025-ND-CP", [1, 6, 7, 8, 30]),
    "ND83-2026_toan_van": ("ND83-2026-ND-CP", None),
    "TT96-2020_Dieu10_PhuLucIV_muc6": ("TT96-2020-TT-BTC", [1, 8, 9, 47, 54, 55, 56, 57]),
}
REPORT_PAGES = {
    "HPG_BCPTBV_2025_trang_9-33-35-39-104-105": ("HPG_Sustainability_Report_2025.pdf",
                                                 [9, 33, 35, 39, 104, 105]),
    "HPG_BCTN_2024_trang_7-23-46": ("HPG_Annual_Report_2024.pdf", [7, 23, 46]),
}


def fold(text: str) -> str:
    text = text.replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in unicodedata.normalize("NFD", text)
                   if unicodedata.category(c) != "Mn").lower()


def header(ws, values: list[str], widths: list[int]) -> None:
    ws.append(values)
    # Slice to the header: a merged note above can make the row wider than it.
    for c, w in zip(ws[ws.max_row][:len(values)], widths, strict=True):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = WRAP
        ws.column_dimensions[c.column_letter].width = w
    # A coordinate string, not ws.cell(): touching a cell creates it, and the
    # next append would then land one row lower than the formulas expect.
    ws.freeze_panes = f"A{ws.max_row + 1}"


def todo(ws, ref: str) -> None:
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"ISBLANK({ref.split(':')[0]})"], fill=TODO))


# ------------------------------------------------------------------ Quỳnh


def hpg_queue(out: Path) -> None:
    run = json.loads(HPG_RUN.read_text(encoding="utf-8"))
    claims = {c["claim_id"]: c for c in run["claims"]}
    verdicts = {v["claim"]["claim_id"]: v for v in run["verifications"]}
    figures = {f["figure_id"]: f for f in run["disclosed_figures"]}
    queue = sorted((p for p in run["priorities"] if p["in_queue"]), key=lambda p: p["rank"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Hàng đợi HPG"
    dv = lp.dropdowns(wb, {
        "yn": ("A", ["Có", "Không", "Không chắc"]),
        "why": ("B", ["giải thích quy trình / kiến thức chung", "tiêu đề, mục lục, dòng bảng",
                      "câu bị cắt hoặc ghép lỗi", "nói về ngành / quốc gia, không phải DN",
                      "chính sách chung, không có nội dung môi trường cụ thể", "khác (ghi lý do)"]),
    }, sheet="DS")
    lists = wb.create_sheet("DS")
    for i, v in enumerate(["Có", "Không", "Không chắc"], 1):
        lists[f"A{i}"] = v
    for i, v in enumerate(["giải thích quy trình / kiến thức chung", "tiêu đề, mục lục, dòng bảng",
                           "câu bị cắt hoặc ghép lỗi", "nói về ngành / quốc gia, không phải DN",
                           "chính sách chung, không có nội dung môi trường cụ thể",
                           "khác (ghi lý do)"], 1):
        lists[f"B{i}"] = v
    lists.sheet_state = "hidden"
    for v in dv.values():
        ws.add_data_validation(v)

    ws.append(["Soát tay hàng đợi HPG — run 185480a0c8b747c4 (25/09/2026, BCPTBV 2025 + BCTN 2024)"])
    ws["A1"].font = Font(bold=True, size=13, color=GREEN)
    ws.append(["Mỗi dòng trả lời MỘT câu: đây có thật là tuyên bố môi trường của chính Hòa Phát không? "
               "Mục 1–14 bắt buộc; 15–22 làm thêm nếu còn thời gian. Mở PDF trong thư mục "
               "06_Bao_cao_HPG đúng trang nếu câu bị cắt khó hiểu."])
    ws.merge_cells("A2:L2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 34
    header(ws, ["#", "Bắt buộc", "Loại mục", "Tài liệu", "Trang", "Loại (hệ thống)",
                "Kết luận hệ thống", "Nội dung hệ thống đưa lên hàng đợi",
                "Có phải tuyên bố MT của chính DN?", "Nếu Không: vì sao?", "Một dòng lý do", "item_id"],
           [5, 9, 10, 16, 7, 18, 20, 80, 16, 30, 40, 4])
    first = ws.max_row + 1
    for p in queue:
        if p["item_type"] == "claim":
            c = claims[p["claim_id"]]
            doc, page, ctype = c["source_name"], c["source_page"], c["claim_type"]
            status = verdicts[c["claim_id"]]["status"]
            kind = "tuyên bố"
        else:
            f = figures.get(p["item_id"], {})
            doc, page, ctype, status, kind = f.get("source_name", ""), f.get("source_page"), \
                f.get("metric") or "", "", "số liệu"
        short = "BCPTBV 2025" if "Sustainability" in doc else "BCTN 2024" if "Annual" in doc else doc
        ws.append([p["rank"], "bắt buộc" if p["rank"] <= 14 else "thêm", kind, short, page, ctype,
                   status, lp.clean(p["text"]), None, None, None, p["item_id"]])
        r = ws.max_row
        for c in ws[r]:
            c.alignment = WRAP
        ws.row_dimensions[r].height = min(160, 16 * (len(p["text"]) // 75 + 2))
        dv["yn"].add(f"I{r}")
        dv["why"].add(f"J{r}")
    last = ws.max_row
    todo(ws, f"I{first}:I{first + 13}")
    ws.column_dimensions["L"].hidden = True

    ws.append([])
    ws.append(["", "", "", "", "", "", "Tự tính:", "Tỷ lệ 'Có' trong 14 mục đầu",
               f'=IF(COUNTA(I{first}:I{first + 13})=0,"",COUNTIF(I{first}:I{first + 13},"Có")'
               f'/COUNTA(I{first}:I{first + 13}))'])
    ws.append(["", "", "", "", "", "", "", "Tỷ lệ 'Có' trên mọi mục đã trả lời",
               f'=IF(COUNTA(I{first}:I{last})=0,"",COUNTIF(I{first}:I{last},"Có")/COUNTA(I{first}:I{last}))'])
    for r in (last + 2, last + 3):
        ws.cell(r, 9).number_format = "0%"
        ws.cell(r, 8).font = Font(bold=True)
    wb.save(out)


# Chosen by hand from the 136 HPG claims outside the review queue: an automatic
# pick drew extraction debris ("Xây dựng và hạ tầng 52% …", the chairman's
# foreword, a World Steel statistic), and timing a search for evidence of a
# sentence that is not a claim measures nothing. 11 carry a figure, 9 are
# commitments or actions; 52f0c981 ("2025 is the first group-wide GHG report")
# sits against BCTN 2024 p.81 ("completed group-wide GHG inventory") on purpose.
TIMING_CLAIMS = [
    "7c135c6baa0e7a15", "e01c2e5e671406b3", "f58847520a3f3c51", "14020e07c47e25c8",
    "b1fd2aa00ea3e6f0", "3a2a1f504958fcaf", "52f0c981b08d6667", "a5c82f7543ddeed2",
    "0f6e1196e910c53d", "743cafdbeee79822", "a7d6165c392417aa", "70d90132a36d7896",
    "5fa78e6aa6349ce6", "3cf6ede544f700f5", "61da3d96da32cfa7", "4b3337ea2123b73e",
    "01ca2d498bb7f5a9", "b62878f07c629a3e", "ad62bba1ae0e923f", "0d779de7787b26b3",
]


def timing_study(out: Path) -> None:
    run = json.loads(HPG_RUN.read_text(encoding="utf-8"))
    claims = {c["claim_id"]: c for c in run["claims"]}
    picked = [claims[cid] for cid in TIMING_CLAIMS]
    # Alternate down a list ordered by type and page: 10/10, the same mix of
    # topics on both sides, and near-duplicates (two rooftop-solar claims) split.
    ordered = sorted(picked, key=lambda c: (c["claim_type"], c["source_page"] or 0))
    groups = {c["claim_id"]: "A · làm tay" if i % 2 == 0 else "B · GreenScan"
              for i, c in enumerate(ordered)}
    rng = random.Random(28)
    rng.shuffle(picked)

    wb = Workbook()
    ws = wb.active
    ws.title = "Bài đo"
    lists = wb.create_sheet("DS")
    verdict_codes = [v.split(" — ")[0] for v in lp.VERDICTS]
    for i, v in enumerate(verdict_codes, 1):
        lists[f"A{i}"] = v
    for i, v in enumerate(["1", "2", "3", "4", "5"], 1):
        lists[f"B{i}"] = v
    lists.sheet_state = "hidden"
    dv = lp.dropdowns(wb, {"verdict": ("A", verdict_codes), "conf": ("B", list("12345"))})
    for v in dv.values():
        ws.add_data_validation(v)

    ws.append(["Bài đo thời gian — 20 tuyên bố HPG, hai nhóm"])
    ws["A1"].font = Font(bold=True, size=13, color=GREEN)
    ws.append(["Nhóm A: làm tay hoàn toàn (mở PDF, Ctrl+F). Nhóm B: dùng GreenScan (Nhân mở sẵn). "
               "Mỗi câu làm MỘT lần, theo đúng nhóm. Gõ giờ dạng 14:05. "
               "Không làm cùng một câu hai lần — lần thứ hai nhanh hơn vì đã nhớ, số đo sẽ sai."])
    ws.merge_cells("A2:M2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 34
    header(ws, ["#", "Nhóm", "Tuyên bố", "Tài liệu", "Trang", "Bắt đầu", "Kết thúc", "Số phút",
                "Trang bằng chứng tìm được (vd: 35; 104)", "Số đoạn bằng chứng", "Kết luận của bạn",
                "Tự tin (1–5)", "Ghi chú", "Số đoạn ĐÚNG (điền sau khi soát chung)"],
           [4, 14, 70, 13, 7, 9, 9, 9, 22, 11, 24, 9, 30, 16])
    first = ws.max_row + 1
    for i, c in enumerate(sorted(picked, key=lambda c: groups[c["claim_id"]]), 1):
        doc = "BCPTBV 2025" if "Sustainability" in c["source_name"] else "BCTN 2024"
        r = ws.max_row + 1
        ws.append([i, groups[c["claim_id"]], lp.clean(c["text"]), doc, c["source_page"], None, None,
                   f'=IF(OR(F{r}="",G{r}=""),"",ROUND((G{r}-F{r})*1440,1))',
                   None, None, None, None, None, None])
        for cell in ws[r]:
            cell.alignment = WRAP
        ws.row_dimensions[r].height = 16 * (len(c["text"]) // 65 + 2)
        ws.cell(r, 6).number_format = "hh:mm"
        ws.cell(r, 7).number_format = "hh:mm"
        dv["verdict"].add(f"K{r}")
        dv["conf"].add(f"L{r}")
    last = ws.max_row
    todo(ws, f"F{first}:G{last}")
    ws.append([])
    for label, grp in (("Trung bình số phút — nhóm A (làm tay)", "A · làm tay"),
                       ("Trung bình số phút — nhóm B (GreenScan)", "B · GreenScan")):
        ws.append(["", "", label, "", "", "", "",
                   f'=IFERROR(AVERAGEIF(B{first}:B{last},"{grp}",H{first}:H{last}),"")'])
        ws.cell(ws.max_row, 3).font = Font(bold=True)
    wb.save(out)


# ------------------------------------------------------------------- Thảo


def _extract(src_path: Path, pages: list[int] | None, label: str, out: Path) -> None:
    """Copy pages unchanged and stamp each with the page number of the original file.

    Printed page numbers differ from file pages (QĐ 13 prints "33" on file page
    37), and every reference in the pack uses file pages, so each copied page says
    which one it is, and the bookmarks list them.
    """
    src = fitz.open(src_path)
    dst = fitz.open()
    pages = list(pages or range(1, src.page_count + 1))
    toc = []
    for n, p in enumerate(pages, 1):
        dst.insert_pdf(src, from_page=p - 1, to_page=p - 1)
        page = dst[-1]
        page.insert_text((18, 14), f"{label} — trang PDF gốc {p}", fontsize=8,
                         fontname="arial", fontfile=STAMP_FONT, color=(0.75, 0.1, 0.1))
        toc.append([1, f"Trang PDF gốc {p}", n])
    dst.set_toc(toc)
    dst.subset_fonts()
    dst.save(out, garbage=3, deflate=True)


def legal_extracts(folder: Path) -> list[tuple[str, str, list[int] | None]]:
    folder.mkdir(parents=True, exist_ok=True)
    made = []
    for name, (doc_id, pages) in LEGAL_PAGES.items():
        _extract(LEGAL / doc_id / "original.pdf", pages, doc_id, folder / f"{name}.pdf")
        made.append((name, doc_id, pages))
    return made


def report_extracts(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for name, (file, pages) in REPORT_PAGES.items():
        _extract(HPG_PDF / file, pages, file, folder / f"{name}.pdf")


def _ocr_pages(doc_id: str) -> list[str]:
    meta = json.loads((LEGAL / doc_id / "ocr_metadata.json").read_text(encoding="utf-8"))
    text = (LEGAL / doc_id / "ocr.txt").read_text(encoding="utf-8")
    pages, pos = [], 0
    for p in meta["per_page"]:
        pages.append(text[pos:pos + p["chars"]])
        pos += p["chars"] + 2   # tools/ocr_legal_document.py joins pages with a blank line
    return pages


QD13_NAMES = {
    "HPG": ["hoa phat"], "HSG": ["hoa sen"], "NKG": ["nam kim"], "VNM": ["sua viet nam"],
    "SAB": ["bia sai gon", "bia - ruou - nuoc giai khat sai gon"], "HT1": ["xi mang ha tien"],
    "DGC": ["hoa chat duc giang", "hoa chat duc giang"], "DPM": ["dam phu my"],
    "DCM": ["dam ca mau", "phan bon dau khi ca mau"], "MSN": ["masan"],
    "BSR": ["loc hoa dau binh son"], "POW": ["dien luc dau khi"], "BMP": ["nhua binh minh"],
    "AAA": ["an phat"], "TCM": ["thanh cong"], "DHG": ["duoc hau giang"], "TRA": ["traphaco"],
    "DPR": ["cao su dong phu"],
}


def qd13_lines() -> list[tuple[str, int, str]]:
    hits = []
    for page_no, page in enumerate(_ocr_pages("QD13-2024-QD-TTg"), 1):
        if page_no < 5:
            continue
        for line in page.splitlines():
            low = fold(line)
            for ticker, names in QD13_NAMES.items():
                if any(n in low for n in names):
                    hits.append((ticker, page_no, " ".join(line.split())))
                    break
    return hits


RULE_TERMS = {
    "QT1 · QĐ 13 — có kiểm kê KNK không": ["kiểm kê khí nhà kính", "kiểm kê knk", "iso 14064", "13/2024"],
    "QT2 · Kỳ báo cáo kiểm kê": ["báo cáo kiểm kê", "hai năm một lần", "06/2022", "119/2025",
                                 "31 tháng 3", "31/3"],
    "QT3 · TT 96 mục 6": ["tổng phát thải khí nhà kính", "phạm vi 1", "tiêu thụ nước", "xử phạt",
                          "tiêu thụ năng lượng"],
}
# A page must also mention this to count: "31/3" alone is any date in a report.
RULE_CONTEXT = {"QT2 · Kỳ báo cáo kiểm kê": "kiểm kê"}
TT96_FIELDS = {
    "6.1 KNK": [r"phát thải khí nhà kính", r"phạm vi 1", r"scope 1", r"tco2"],
    "6.2 NVL": [r"nguyên vật liệu"],
    "6.3 Năng lượng": [r"tiêu thụ năng lượng", r"năng lượng tiêu thụ", r"\bkwh\b", r"\bgj\b"],
    "6.4 Nước": [r"tiêu thụ nước", r"nước tiêu thụ", r"lượng nước sử dụng", r"\bm3\b", r"m³"],
    "6.5 Tuân thủ": [r"xử phạt", r"vi phạm.{0,40}môi trường", r"không có vi phạm"],
}
REPORT_TYPES = {"annual_report", "sustainability_report", "integrated_report"}


def _window(text: str, term: str, width: int = 260) -> str:
    low = text.lower()
    i = low.find(term)
    start, end = max(0, i - width), min(len(text), i + len(term) + width)
    return ("…" if start else "") + lp.clean(text[start:end]) + ("…" if end < len(text) else "")


def report_pages() -> list[dict]:
    """Every page of every report the team has: corpus chunks plus the two HPG PDFs."""
    docs = {}
    for line in (CRAWL / "documents.jsonl").read_text(encoding="utf-8").splitlines():
        d = json.loads(line)
        docs[d["document_id"]] = d
    pages = []
    for line in (CRAWL / "chunks.jsonl").read_text(encoding="utf-8").splitlines():
        c = json.loads(line)
        if c["document_type"] not in REPORT_TYPES:
            continue
        d = docs.get(c["document_id"], {})
        pages.append({"ticker": c["ticker"], "doc_id": c["document_id"],
                      "doc": lp.DOC_TYPES.get(c["document_type"], c["document_type"]),
                      "year": c.get("publication_year") or "?", "page": c["unit_index"],
                      "title": lp.clean(d.get("title") or "")[:60], "url": d.get("url") or "",
                      "text": c["text"]})
    for file, doc, year in (("HPG_Sustainability_Report_2025.pdf", "BC PTBV", 2025),
                            ("HPG_Annual_Report_2024.pdf", "BC thường niên", 2024)):
        for i, page in enumerate(fitz.open(HPG_PDF / file), 1):
            pages.append({"ticker": "HPG", "doc_id": file, "doc": doc, "year": year, "page": i,
                          "title": file, "url": f"Goi_Thao/07_Bao_cao_HPG_trich hoặc bản đầy đủ: {file}",
                          "text": page.get_text()})
    return pages


def candidates_workbook(out: Path) -> None:
    pages = report_pages()
    wb = Workbook()
    lists = wb.create_sheet("DS")
    for i, v in enumerate(["đạt", "không đạt", "miễn trừ (không áp dụng)", "không dùng"], 1):
        lists[f"A{i}"] = v
    for i, v in enumerate(["khớp đúng DN", "không phải DN này", "không chắc"], 1):
        lists[f"B{i}"] = v
    lists.sheet_state = "hidden"
    dv = lp.dropdowns(wb, {"pick": ("A", ["đạt", "không đạt", "miễn trừ (không áp dụng)", "không dùng"]),
                           "match": ("B", ["khớp đúng DN", "không phải DN này", "không chắc"])})

    ws = wb.active
    ws.title = "Trích báo cáo"
    ws.add_data_validation(dv["pick"])
    header(ws, ["Quy tắc", "DN", "Tài liệu", "Năm", "Trang", "Từ khoá gặp", "Trích đoạn (±260 ký tự)",
                "Thảo chọn làm ví dụ", "Ghi chú", "Nguồn (URL / tệp)"],
           [22, 7, 15, 7, 7, 16, 90, 18, 30, 40])
    for rule, terms in RULE_TERMS.items():
        per_doc = Counter()
        rows = []
        # HPG first: it is the demo company and its pages are in the pack.
        for p in sorted(pages, key=lambda p: (p["ticker"] != "HPG", p["ticker"], str(p["year"]), p["page"])):
            low = p["text"].lower()
            term = next((t for t in terms if t in low), None)
            if rule in RULE_CONTEXT and RULE_CONTEXT[rule] not in low:
                term = None
            if not term or per_doc[p["doc_id"]] >= 2:
                continue
            per_doc[p["doc_id"]] += 1
            rows.append([rule, p["ticker"], p["doc"], p["year"], p["page"], term,
                         _window(p["text"], term), None, None, p["url"]])
            if len(rows) >= 24:
                break
        for row in rows:
            ws.append(row)
            r = ws.max_row
            for c in ws[r]:
                c.alignment = WRAP
            ws.row_dimensions[r].height = 105
            dv["pick"].add(f"H{r}")

    ws = wb.create_sheet("TT96 mục 6 — ma trận")
    ws.append(["Trang đầu tiên có từ khoá của từng trường mục 6 TT 96, theo từng báo cáo. "
               "'—' nghĩa là KHÔNG THẤY TỪ KHOÁ, không có nghĩa là báo cáo không công bố "
               "(có thể viết khác chữ, bằng tiếng Anh, hoặc trang là ảnh). Dùng để chọn nhanh ứng viên."])
    ws.merge_cells("A1:K1")
    ws["A1"].alignment = WRAP
    ws.row_dimensions[1].height = 44
    header(ws, ["DN", "Tài liệu", "Năm", "Tiêu đề / tệp", *TT96_FIELDS, "Nguồn"],
           [7, 15, 7, 40, 10, 10, 13, 10, 12, 40])
    by_doc = defaultdict(list)
    for p in pages:
        by_doc[p["doc_id"]].append(p)
    rows = []
    for doc_pages in by_doc.values():
        first = {}
        for field, pats in TT96_FIELDS.items():
            hits = [p["page"] for p in doc_pages if any(re.search(x, p["text"].lower()) for x in pats)]
            first[field] = min(hits) if hits else "—"
        if all(v == "—" for v in first.values()):
            continue
        p0 = doc_pages[0]
        rows.append([p0["ticker"], p0["doc"], p0["year"], p0["title"], *first.values(), p0["url"]])
    for row in sorted(rows, key=lambda r: (r[0] != "HPG", r[0], str(r[2]))):
        ws.append(row)

    ws = wb.create_sheet("QĐ13 — cơ sở DN niêm yết")
    ws.add_data_validation(dv["match"])
    ws.append(["Dòng trong Phụ lục II–V QĐ 13/2024 (bản OCR) có tên giống cơ sở của DN niêm yết. "
               "OCR nhiễu và bảng bị xuống dòng, nên MỌI dòng phải đối chiếu với trang gốc "
               "(03_Van_ban_goc_trich/QD13-2024_trang_co_so_DN_niem_yet.pdf hoặc bản đầy đủ). "
               "Cột số cuối dòng là TIÊU THỤ NĂNG LƯỢNG (TOE) ở Phụ lục II–IV, công suất xử lý (tấn/năm) "
               "ở Phụ lục V — KHÔNG phải lượng phát thải tCO2e."])
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = WRAP
    ws.row_dimensions[1].height = 58
    header(ws, ["Mã CK (đoán)", "Trang PDF gốc", "Dòng OCR", "Khớp?", "Ghi chú", ""],
           [11, 11, 100, 18, 30, 2])
    for ticker, page, line in qd13_lines():
        ws.append([ticker, page, line, None, None])
        ws.cell(ws.max_row, 3).alignment = WRAP
        dv["match"].add(f"D{ws.max_row}")

    wb.move_sheet(lists, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(lists.title))
    wb.save(out)


LANGUAGE_TERMS = {
    "cấm": ["vi phạm", "gian lận", "sai phạm", "trung thực", "lừa dối", "gian dối", "phạm luật", "phạm pháp"],
    "cân nhắc": ["tẩy xanh", "greenwash", "xếp hạng", "chấm điểm"],
}


def language_review(out: Path) -> None:
    sources = []
    for path in sorted((ROOT / "frontend" / "src").rglob("*")):
        if path.suffix in {".jsx", ".js", ".ts", ".tsx", ".html"}:
            sources.append(("Giao diện (UI)", path.relative_to(ROOT).as_posix(),
                            path.read_text(encoding="utf-8").splitlines()))
    for path in sorted((ROOT / "docs" / "07-presentation").glob("*.md")):
        sources.append(("Q&A / thuyết trình", path.relative_to(ROOT).as_posix(),
                        path.read_text(encoding="utf-8").splitlines()))
    brief = ROOT / "docs" / "07-presentation" / "technical_brief" / "GreenScan_Technical_Brief.tex"
    sources.append(("Tài liệu kỹ thuật", brief.relative_to(ROOT).as_posix(),
                    brief.read_text(encoding="utf-8-sig").splitlines()))
    slides = []
    if SLIDES.exists():
        with zipfile.ZipFile(SLIDES) as z:
            for name in sorted(z.namelist(), key=lambda s: int(s.split(".")[0])):
                doc = fitz.open(stream=z.read(name), filetype="pdf")
                text = " ".join(p.get_text() for p in doc)
                slides.append((name, " ".join(text.split())))
                sources.append(("Slide", f"GREENSCAN.zip/{name}", [" ".join(text.split())]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Chỗ cần soát"
    lists = wb.create_sheet("DS")
    for i, v in enumerate(["giữ nguyên", "sửa", "xoá"], 1):
        lists[f"A{i}"] = v
    lists.sheet_state = "hidden"
    dv = lp.dropdowns(wb, {"act": ("A", ["giữ nguyên", "sửa", "xoá"])})
    ws.add_data_validation(dv["act"])
    ws.append(["Sản phẩm KHÔNG kết luận doanh nghiệp vi phạm pháp luật — chỉ nêu nghĩa vụ nào chưa được "
               "chứng minh. Mỗi dòng: đọc ngữ cảnh, chọn giữ / sửa / xoá, nếu sửa thì viết câu thay thế. "
               "Câu PHỦ ĐỊNH (\"không có nghĩa là DN vi phạm\") thường giữ được."])
    ws.merge_cells("A1:H1")
    ws["A1"].alignment = WRAP
    ws.row_dimensions[1].height = 44
    header(ws, ["Mức", "Nơi", "Tệp", "Dòng", "Từ", "Ngữ cảnh", "Quyết định", "Câu thay thế / ghi chú"],
           [9, 16, 36, 6, 12, 80, 13, 50])
    for place, file, lines in sources:
        for n, line in enumerate(lines, 1):
            low = line.lower()
            for level, terms in LANGUAGE_TERMS.items():
                for term in terms:
                    i = low.find(term)
                    if i < 0:
                        continue
                    start = max(0, i - 160)
                    context = ("…" if start else "") + line[start:i + 200].strip() + "…"
                    ws.append([level, place, file, n if place != "Slide" else "", term, context])
                    r = ws.max_row
                    for c in ws[r]:
                        c.alignment = WRAP
                    ws.row_dimensions[r].height = 62
                    dv["act"].add(f"G{r}")
    ws2 = wb.create_sheet("Toàn văn slide")
    header(ws2, ["Slide", "Chữ trên slide (slide chỉ có ảnh thì trống — mở PDF để xem)", "Nhận xét của Thảo"],
           [10, 110, 40])
    for name, text in slides:
        ws2.append([name, text or "(slide là ảnh — không trích được chữ)"])
        ws2.cell(ws2.max_row, 2).alignment = WRAP
        ws2.row_dimensions[ws2.max_row].height = max(30, 15 * (len(text) // 105 + 1))
    wb.move_sheet(lists, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(lists.title))
    wb.save(out)


# ---------------------------------------------------------------- Word form


def _shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    props.append(shd)


RULES = [
    {
        "id": "VN_GHG_INVENTORY_001",
        "title": "Quy tắc 1 — Cơ sở thuộc danh mục phải kiểm kê khí nhà kính (QĐ 13/2024/QĐ-TTg)",
        "question": "Doanh nghiệp có cơ sở nằm trong danh mục phải kiểm kê KNK không? Nếu có, báo cáo "
                    "có nói đã kiểm kê KNK cấp cơ sở không?",
        "q1": "Điều kiện 1 — facility_in_list: ít nhất một cơ sở của DN có tên trong Phụ lục II–V. "
              "Cách kiểm: NGƯỜI XEM (manual) — tên trong phụ lục là tên cơ sở/công ty con, OCR nhiễu; "
              "Nhân dựng bảng đối chiếu cơ sở ↔ mã CK, người duyệt.\n"
              "Điều kiện 2 — inventory_disclosed: báo cáo nói đã kiểm kê KNK cấp cơ sở. Cách kiểm: "
              "TỪ KHOÁ: 'kiểm kê khí nhà kính', 'kiểm kê KNK', 'ISO 14064'.",
        "q2": "QĐ 13/2024/QĐ-TTg: Điều 1 khoản 1–5 (Phụ lục I lĩnh vực; II Công Thương; III GTVT; "
              "IV Xây dựng; V TN&MT) — trang 1. Điều 2 khoản 1: cơ sở thực hiện kiểm kê cấp cơ sở, nộp "
              "báo cáo theo NĐ 06/2022 — trang 1–2.",
        "q3": "Ký 13/08/2024; HIỆU LỰC 01/10/2024 (Điều 3 khoản 1, trang 2) — mẫu YAML trong bàn giao "
              "25/09 ghi 2024-08-13 là ngày ký, không phải ngày hiệu lực. Thay thế QĐ 01/2022/QĐ-TTg. "
              "Điều 3 khoản 2: cơ sở có trong QĐ 01/2022 nhưng không có trong QĐ 13 không phải nộp báo "
              "cáo cấp cơ sở năm 2025.",
        "q4": "Ngành: theo Phụ lục I (năng lượng, GTVT, xây dựng, quá trình công nghiệp: hoá chất, luyện "
              "kim…, nông lâm nghiệp, chất thải). Đối tượng: nghĩa vụ đặt lên CƠ SỞ (facility). "
              "CẢNH BÁO: báo cáo DN niêm yết là cấp tập đoàn/công ty mẹ — nếu chỉ ghi [facility], hệ "
              "thống sẽ bỏ qua báo cáo tập đoàn và MẤT finding. Đề xuất: để trống đối tượng, hoặc ghi "
              "[facility, company, group].",
        "q5": "ĐẠT (ứng viên): HPG BCPTBV 2025 trang 35 — \"Tất cả các công ty thành viên của Tập đoàn "
              "thuộc danh mục phải kiểm kê phát thải theo Nghị định 06/2022/NĐ-CP đã tiến hành kiểm kê "
              "KNK theo chuẩn ISO 14064 … nộp báo cáo kiểm kê KNK đầy đủ cho các cơ quan chức năng từ "
              "năm 2024.\" Cơ sở HPG có trong Phụ lục II (trang PDF 37, 45–47, 66, 68).\n"
              "KHÔNG ĐẠT: chưa có — chọn trong 04_Ung_vien_vi_du.xlsx: một DN có cơ sở trong QĐ 13 "
              "mà báo cáo cùng năm không nhắc kiểm kê.",
    },
    {
        "id": "VN_GHG_REPORTING_PERIOD_001",
        "title": "Quy tắc 2 — Kỳ và hạn nộp báo cáo kiểm kê KNK cấp cơ sở (NĐ 06/2022 → 119/2025 → 83/2026)",
        "question": "Điều báo cáo nói về kỳ kiểm kê / thời điểm nộp có khớp với quy định có hiệu lực tại "
                    "ngày công bố không?",
        "q1": "Điều kiện 1 — cycle_stated: báo cáo nêu năm kiểm kê hoặc năm nộp. Cách kiểm: TỪ KHOÁ "
              "('báo cáo kiểm kê', 'nộp báo cáo', 'hai năm một lần') để tìm đoạn, còn khớp kỳ là NGƯỜI XEM.\n"
              "Điều kiện 2 — cycle_matches: năm/kỳ nêu trong báo cáo đúng với quy định tại ngày công bố. "
              "Cách kiểm: NGƯỜI XEM (manual) — so năm là suy luận pháp lý, không phải khớp chữ.",
        "q2": "NĐ 06/2022/NĐ-CP Điều 11 khoản 4: (a) cung cấp số liệu hoạt động năm trước, trước 31/3 "
              "từ 2023; (b) kiểm kê cấp cơ sở, báo cáo định kỳ HAI NĂM MỘT LẦN cho năm 2024 trở đi (Mẫu 06 "
              "Phụ lục II), gửi UBND cấp tỉnh trước 31/3 từ 2025; (c) hoàn thiện, gửi Bộ trước 01/12 của "
              "kỳ báo cáo từ 2025 — trang 9–11.\n"
              "NĐ 119/2025/NĐ-CP Điều 1 khoản 7 (sửa Điều 11): bổ sung điểm e khoản 1 (báo cáo hai năm "
              "một lần gồm kết quả hai năm liền kề năm nộp); điểm c khoản 4: NHIỆT ĐIỆN, SẮT THÉP, XI MĂNG "
              "trong danh mục báo cáo hai năm một lần cho năm 2026 trở đi; điểm d: cơ sở được phân bổ hạn "
              "ngạch từ 2027 cho năm 2028 trở đi; khoản 6a: gửi báo cáo đã thẩm định trước 01/12 từ 2027 "
              "— trang 6–8.",
        "q3": "NĐ 06/2022 hiệu lực 07/01/2022. NĐ 119/2025 ký 09/06/2025, HIỆU LỰC 01/08/2025 (Điều 2, "
              "trang 30) — hệ thống đang ghi 09/06/2025, cần sửa. NĐ 83/2026 (hiệu lực từ ngày ký "
              "01/03/2026): theo bản OCR chỉ sửa quy định về CHẤT ĐƯỢC KIỂM SOÁT (HCFC, tầng ô-dôn), có vẻ "
              "KHÔNG đổi kỳ kiểm kê — Thảo xác nhận; nếu đúng thì không được dẫn 83/2026 làm căn cứ cho "
              "tuyên bố phát thải (hệ thống đang dẫn).",
        "q4": "Ngành: mọi cơ sở trong danh mục QĐ 13; riêng nhiệt điện / sắt thép / xi măng có lịch riêng "
              "(điểm c) → có thể cần tách thành 2 quy tắc: 2a [power, steel, cement] và 2b các ngành "
              "còn lại. Đối tượng: cơ sở (xem cảnh báo ở quy tắc 1).",
        "q5": "ĐẠT (ứng viên): HPG BCTN 2024 trang 23 — \"… nộp báo cáo kiểm kê khí nhà kính cấp cơ sở về "
              "cơ quan có thẩm quyền từ tháng 3/2025\" (khớp hạn 31/3 từ 2025).\n"
              "CẦN THẢO ĐÁNH GIÁ: HPG BCPTBV 2025 trang 35 — \"… nộp báo cáo kiểm kê KNK đầy đủ … từ năm "
              "2024\" (nộp năm 2024 hay kiểm kê cho năm 2024?).\n"
              "KHÔNG ĐẠT: chưa có — chọn trong 04_Ung_vien_vi_du.xlsx.",
    },
    {
        "id": "VN_DISCLOSURE_TT96_ENV_001",
        "title": "Quy tắc 3 — Trường môi trường bắt buộc trong báo cáo thường niên (TT 96/2020, Phụ lục IV mục 6)",
        "question": "Báo cáo thường niên (hoặc BCPTBV lập riêng thay cho mục 6) có đủ các trường môi "
                    "trường bắt buộc không?",
        "q1": "Mỗi trường là một điều kiện, kiểm bằng TỪ KHOÁ trong mục môi trường:\n"
              "6.1 KNK: 'phát thải khí nhà kính', 'Phạm vi 1', 'Scope 1', 'tCO2' · "
              "6.2 NVL: 'nguyên vật liệu', 'tái chế' · 6.3 Năng lượng: 'tiêu thụ năng lượng', 'kWh', 'GJ' · "
              "6.4 Nước: 'tiêu thụ nước', 'm3' · 6.5 Tuân thủ: 'xử phạt', 'không có vi phạm'.\n"
              "Lưu ý: từ khoá chỉ tìm được CÓ; không tìm thấy thì hệ thống trả UNKNOWN (chưa biết), "
              "không trả KHÔNG ĐẠT.",
        "q2": "TT 96/2020/TT-BTC Điều 10 khoản 2: lập BCTN theo mẫu Phụ lục IV, công bố trong 20 ngày từ "
              "ngày công bố BCTC năm kiểm toán, không quá 110 ngày từ khi kết thúc năm tài chính — "
              "trang 8–9. Phụ lục IV, mục 6 \"Báo cáo tác động liên quan đến môi trường và xã hội\": "
              "6.1 tổng phát thải KNK trực tiếp và gián tiếp + sáng kiến giảm; 6.2 NVL (tổng lượng; % tái "
              "chế); 6.3 năng lượng (trực tiếp/gián tiếp; tiết kiệm; sáng kiến); 6.4 nước (nguồn và lượng; "
              "% tái chế/tái sử dụng); 6.5 tuân thủ pháp luật BVMT (số lần bị xử phạt; tổng tiền phạt); "
              "6.6–6.8 lao động, cộng đồng, vốn xanh — trang 54–56. Phần đánh giá của Ban Giám đốc mục 6 "
              "(chỉ tiêu môi trường) — trang 57.",
        "q3": "Hiệu lực 01/01/2021. Hệ thống ghi \"có văn bản sửa đổi\" nhưng CHƯA ghi văn bản nào — "
              "Thảo tra và điền (số hiệu, ngày hiệu lực, có đụng tới Phụ lục IV mục 6 không).",
        "q4": "Đối tượng: công ty đại chúng (Điều 2 khoản 1 điểm a). LƯU Ý QUAN TRỌNG (Phụ lục IV, ghi chú "
              "trang 56): mục 6 có thể lập riêng thành BCPTBV, và 6.1, 6.2, 6.3 KHÔNG BẮT BUỘC với DN "
              "dịch vụ tài chính, ngân hàng, chứng khoán, bảo hiểm. → Đề xuất tách: 3a (6.1–6.3) áp cho "
              "mọi ngành TRỪ tài chính/ngân hàng/chứng khoán/bảo hiểm; 3b (6.4–6.5) áp cho mọi công ty "
              "đại chúng. Hệ thống hiện chỉ có danh sách \"áp cho\", chưa có \"trừ\" — ghi rõ ngành bị "
              "trừ, Nhân xử lý kỹ thuật.",
        "q5": "ĐẠT (ứng viên): HPG BCPTBV 2025 trang 104 — Phạm vi 1: 22.540.603 tấn CO2e; Phạm vi 2: "
              "933.876 tấn; năng lượng 193.403.521 GJ; trang 105 — nước sử dụng 41.334 triệu lít.\n"
              "MIỄN TRỪ (ví dụ cho 3a): BVH (bảo hiểm) — 6.1–6.3 không bắt buộc.\n"
              "KHÔNG ĐẠT (ứng viên): xem sheet \"TT96 mục 6 — ma trận\" — vd. DPR BCTN 2024 có 6.3/6.4/6.5 "
              "nhưng không thấy từ khoá 6.1 (DN cao su, không thuộc diện miễn). Cần đọc trang gốc.",
    },
]
SECTORS = ("steel (thép) · cement (xi măng) · power (điện) · chemicals (hoá chất) · fertilizer (phân bón) · "
           "oil_gas (dầu khí) · textiles (dệt may) · food_beverage (thực phẩm, đồ uống) · "
           "plastics_paper (nhựa, giấy, bao bì) · industrial_parks (BĐS KCN) · logistics (vận tải) · "
           "aviation (hàng không) · construction (xây dựng) · agriculture (nông nghiệp) · "
           "pharma (dược) · retail (bán lẻ) · banking (ngân hàng) · insurance (bảo hiểm) · "
           "securities (chứng khoán) · finance_other (tài chính khác)")
SUBJECTS = ("facility (cơ sở) · company (doanh nghiệp / pháp nhân) · group (tập đoàn, số hợp nhất) · "
            "public_company (công ty đại chúng) · credit_institution (tổ chức tín dụng)")


def rule_form(out: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)

    title = doc.add_heading("Phiếu 3 quy tắc pháp lý — Thảo điền", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(
        "Cột giữa là ĐỀ XUẤT SẴN do Nhân soạn từ bản OCR — có thể sai. Việc của Thảo: đối chiếu từng ô "
        "với trang gốc trong thư mục 03_Van_ban_goc_trich, rồi ở cột phải ghi \"Đúng\" hoặc viết lại. "
        "Một quy tắc chỉ được code khi đủ cả 5 câu. Không chắc thì ghi \"chưa chắc\" + lý do — không đoán.")
    p = doc.add_paragraph()
    p.add_run("Hai quy tắc vàng. ").bold = True
    p.add_run("(1) Để trống ngành/đối tượng nghĩa là \"áp cho mọi đối tượng\" — chỉ điền khi chắc chắn, "
              "vì điền sai sẽ LÀM MẤT finding đúng. (2) Điều kiện không kiểm được bằng từ khoá thì ghi "
              "\"người xem\" — hệ thống sẽ trả \"chưa biết\" và chuyển người, không bao giờ trả \"không đạt\". "
              "Chưa biết ≠ không đáp ứng.")

    rows = [("Câu hỏi quy tắc trả lời", "question"),
            ("① Đầu vào trích được tự động từ PDF? (từng điều kiện: mô tả · cách kiểm · từ khoá)", "q1"),
            ("② Điều / khoản nguồn chính thức (số hiệu + điều + khoản + trang)", "q2"),
            ("③ Ngày hiệu lực; văn bản sửa đổi / thay thế", "q3"),
            ("④ Áp cho ngành nào, đối tượng nào (và ngành nào bị trừ)", "q4"),
            ("⑤ Một ví dụ ĐẠT và một ví dụ KHÔNG ĐẠT (DN · tài liệu · năm · trang · trích nguyên văn)", "q5")]
    for rule in RULES:
        doc.add_heading(rule["title"], level=1)
        doc.add_paragraph(f"Mã đề xuất: {rule['id']}").runs[0].italic = True
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        widths = (Cm(4.2), Cm(8.3), Cm(5.0))
        for cell, text, w in zip(table.rows[0].cells, ("Câu hỏi", "Đề xuất sẵn (cần kiểm)",
                                                          "Thảo xác nhận / sửa"), widths, strict=True):
            cell.text = text
            cell.width = w
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, GREEN)
        for label, key in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[0].paragraphs[0].runs[0].bold = True
            cells[1].text = rule[key]
            for cell, w in zip(cells, widths, strict=True):
                cell.width = w
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9.5)
            _shade(cells[2], "FFF4CC")
        cells = table.add_row().cells
        cells[0].text = "Người soát · ngày"
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = "reviewed_by / reviewed_at — bắt buộc trước khi quy tắc vào sản phẩm"
        _shade(cells[2], "FFF4CC")
        doc.add_paragraph()

    doc.add_heading("Từ vựng cho câu ④ — chọn đúng chữ trong danh sách", level=1)
    doc.add_paragraph("Hệ thống so khớp NGUYÊN CHỮ, nên phải dùng đúng mã tiếng Anh bên dưới "
                      "(thiếu ngành nào thì ghi thêm, Nhân bổ sung vào danh sách).")
    p = doc.add_paragraph()
    p.add_run("Ngành: ").bold = True
    p.add_run(SECTORS)
    p = doc.add_paragraph()
    p.add_run("Đối tượng: ").bold = True
    p.add_run(SUBJECTS)

    doc.add_heading("Phần của Nhân (Thảo không cần điền)", level=1)
    doc.add_paragraph(
        "legal_issue · claim_types · check (terms_present / numeric_threshold / manual) · "
        "source_clauses dạng mã (vd. QD13-2024-QD-TTg:article_1:paragraph_2) · "
        "machine_executable · unit test. Lưu ý kỹ thuật: tuyên bố emissions_reduction được hệ thống "
        "quy về vấn đề 'emissions' — quy tắc khai legal_issue 'ghg_inventory' như mẫu bàn giao 25/09 "
        "sẽ không bao giờ được chọn.")
    doc.save(out)


# -------------------------------------------------------------------- build


def build() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    for folder in (QUYNH, THAO):
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)

    for session, stem in LOTS:
        lp.export(SESSIONS / session, "quynh", QUYNH / f"{stem}_Quynh.xlsx")
    lp.export(SESSIONS / LOTS[0][0], "thao", THAO / f"{LOTS[0][1]}_Thao.xlsx")
    hpg_queue(QUYNH / "04_HPG_Hang_doi_22muc.xlsx")
    timing_study(QUYNH / "05_Bai_do_thoi_gian_20cau.xlsx")
    reports = QUYNH / "06_Bao_cao_HPG"
    reports.mkdir()
    shutil.copy2(HPG_PDF / "HPG_Sustainability_Report_2025.pdf", reports / "HPG_BCPTBV_2025.pdf")
    shutil.copy2(HPG_PDF / "HPG_Annual_Report_2024.pdf", reports / "HPG_BCTN_2024.pdf")

    rule_form(THAO / "02_Phieu_3_quy_tac_phap_ly.docx")
    made = legal_extracts(THAO / "03_Van_ban_goc_trich")
    candidates_workbook(THAO / "04_Ung_vien_vi_du.xlsx")
    language_review(THAO / "05_Soat_ngon_tu.xlsx")
    if SLIDES.exists():
        shutil.copy2(SLIDES, THAO / "06_Slide_GREENSCAN.zip")
    report_extracts(THAO / "07_Bao_cao_HPG_trich")

    for name, doc_id, pages in made:
        print(f"  {name}.pdf  <- {doc_id} trang {pages or 'toàn văn'}")
    for folder in (QUYNH, THAO):
        print(folder.relative_to(ROOT))
        for path in sorted(folder.rglob("*")):
            if path.is_file():
                print(f"   {path.relative_to(folder)}  ({path.stat().st_size // 1024} KB)")


def pack() -> None:
    if not GUIDE.exists():
        raise SystemExit(f"compile the guide first: {GUIDE}")
    for folder, zip_name in ((QUYNH, "Goi_Quynh_2026-09-28.zip"), (THAO, "Goi_Thao_2026-09-28.zip")):
        shutil.copy2(GUIDE, folder / "00_DOC_TRUOC_Huong_dan.pdf")
        with zipfile.ZipFile(OUT / zip_name, "w", zipfile.ZIP_DEFLATED) as z:
            for path in sorted(folder.rglob("*")):
                if path.is_file():
                    z.write(path, Path(folder.name) / path.relative_to(folder))
        print(f"{zip_name}: {(OUT / zip_name).stat().st_size // 1024} KB")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["build", "zip"])
    args = ap.parse_args()
    build() if args.cmd == "build" else pack()


if __name__ == "__main__":
    main()
