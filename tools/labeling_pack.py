#!/usr/bin/env python3
"""Gold labelling in a spreadsheet, for labellers who do not use a terminal.

`label_session.py set` needs one shell command per field, which is the wrong
tool for someone labelling 80 claims. This turns a session into an Excel
workbook with one card per claim (dropdowns, a live completeness check, a
progress sheet), reads the filled workbook back into the session format, and
measures agreement between two labellers.

    python tools/labeling_pack.py export data/gold/sessions/<s>.jsonl --labeler quynh --out X.xlsx
    python tools/labeling_pack.py import X.xlsx            # -> sessions/labelers/<s>_<labeler>.jsonl
    python tools/labeling_pack.py import X.xlsx --final    # -> sessions/<s>.jsonl (adjudicated gold)
    python tools/labeling_pack.py kappa A.jsonl B.jsonl    # -> benchmark/kappa_<s>.md + disagreements .xlsx
    python tools/labeling_pack.py import-all <folder>      # everything sent back: gold, HPG queue, timing, κ

Only `--final` writes where `evaluate_gold.py` reads. A single labeller's copy
is never gold on its own when the batch was labelled twice; it becomes gold
after the two copies are reconciled and one reconciled workbook is imported
with `--final`.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SESSIONS = ROOT / "data" / "gold" / "sessions"
LABELER_COPIES = SESSIONS / "labelers"
BENCHMARK = ROOT / "benchmark"

ATTRS = ["metric", "value_unit", "period", "baseline", "scope"]

# code — Vietnamese gloss. The code before " — " is what is stored.
IS_CLAIM = ["yes — có", "no — không", "unsure — không chắc"]
CLAIM_TYPES = [
    "emissions_reduction — phát thải / KNK",
    "renewable_energy — năng lượng tái tạo",
    "resource_efficiency — tiết kiệm năng lượng, nước, tài nguyên",
    "waste_and_circularity — chất thải, tái chế, tuần hoàn",
    "green_finance — tín dụng / trái phiếu xanh",
    "biodiversity — đa dạng sinh học",
    "esg_process_integration — quy trình ESG của DN",
    "generic_sustainability — chung chung (xanh, bền vững)",
]
ATTR_VALUES = ["present — có đủ", "partial — một phần", "missing — thiếu", "na — không áp dụng"]
RELATIONS = [
    "supports — xác nhận",
    "partial — xác nhận một phần",
    "contradicts — mâu thuẫn",
    "context — bối cảnh",
    "not_relevant — không liên quan",
]
VERDICTS = [
    "SUPPORTED — được hỗ trợ",
    "PARTIALLY_SUPPORTED — hỗ trợ một phần",
    "CONTRADICTED — bị mâu thuẫn",
    "UNSUPPORTED — không được hỗ trợ",
    "INSUFFICIENT_EVIDENCE — chưa đủ bằng chứng",
]
LISTS = {"A": IS_CLAIM, "B": CLAIM_TYPES, "C": ATTR_VALUES, "D": RELATIONS, "E": VERDICTS}

ATTR_QUESTIONS = {
    "metric": ("Chỉ số (metric)",
               "present: gọi tên chỉ số (KNK Scope 1+2, điện, nước…) · partial: chỉ định tính "
               "(\"trung hoà carbon\", \"xanh\") · missing: không có chỉ số"),
    "value_unit": ("Số liệu và đơn vị",
                   "present: có số VÀ đơn vị · partial: có số không đơn vị · missing: không có số · "
                   "na: tuyên bố trạng thái, không cần số (ghi chú)"),
    "period": ("Kỳ / năm",
               "present: năm hoặc kỳ rõ · partial: \"hằng năm\", \"trong 1 năm\" không nói năm nào · "
               "missing: không có mốc thời gian"),
    "baseline": ("Năm gốc (so với khi nào)",
                 "present: năm gốc rõ · partial: \"so với trước\" · missing: có tăng/giảm/mục tiêu "
                 "mà không có năm gốc · na: không phải tăng/giảm"),
    "scope": ("Phạm vi / ranh giới",
              "present: Scope 1/2/3 VÀ ranh giới (toàn tập đoàn / nhà máy) · partial: chỉ một trong "
              "hai · missing: không nói"),
}

# Characters some PDF fonts emit for the "ti"/"fi"/"fl" ligatures.
LIGATURES = {"琀椀": "ti", "昀椀": "fi", "昀氀": "fl", "昀昀": "ff"}
DOC_TYPES = {
    "sustainability_report": "BC PTBV", "annual_report": "BC thường niên",
    "integrated_report": "BC tích hợp", "financial_statement": "BCTC",
    "agm_document": "Tài liệu ĐHĐCĐ", "governance_report": "BC quản trị", "other": "Khác",
}
SPLITS = {"train": "train", "dev": "dev", "test": "test (giữ lại)", "future_holdout": "năm mới nhất"}

GREEN = "1B7F5A"
FILL_HEAD = PatternFill("solid", fgColor=GREEN)
FILL_CLAIM = PatternFill("solid", fgColor="EEF5F1")
FILL_TODO = PatternFill("solid", fgColor="FFF4CC")
FILL_OK = PatternFill("solid", fgColor="E3F4EA")
FILL_WARN = PatternFill("solid", fgColor="FDE2D8")
THIN = Side(style="thin", color="C9D3CD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

DISPLAY_LIMIT = 2200  # characters of evidence shown on the card


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")


def clean(text: str) -> str:
    for bad, good in LIGATURES.items():
        text = text.replace(bad, good)
    return " ".join(text.split())


def code(value) -> str | None:
    """'present — có đủ' -> 'present'. Blank cells are None."""
    if value is None:
        return None
    text = str(value).strip()
    return text.split(" — ")[0].strip() or None if text else None


def _height(text: str, chars_per_line: int) -> float:
    lines = sum(max(1, -(-len(part) // chars_per_line)) for part in text.split("\n"))
    return min(409, max(18, 15.5 * lines + 4))


# ------------------------------------------------------------------- export


def export(session: Path, labeler: str, out: Path) -> Path:
    rows = read_jsonl(session)
    wb = Workbook()

    guide = wb.active
    guide.title = "Đọc trước"
    card = wb.create_sheet("Gán nhãn")
    progress = wb.create_sheet("Tiến độ")
    lists = wb.create_sheet("DS")
    meta = wb.create_sheet("_meta")

    for col, values in LISTS.items():
        for i, v in enumerate(values, start=1):
            lists[f"{col}{i}"] = v
    lists.sheet_state = "hidden"
    meta.append(["session", session.name])
    meta.append(["labeler", labeler])
    meta.append(["exported_at", dt.datetime.now().isoformat(timespec="seconds")])
    meta.append(["tool", "labeling_pack.py v1"])
    meta.sheet_state = "hidden"

    dv = dropdowns(wb, {
        "is_claim": ("A", IS_CLAIM), "claim_type": ("B", CLAIM_TYPES),
        "attr": ("C", ATTR_VALUES), "relation": ("D", RELATIONS), "verdict": ("E", VERDICTS),
    })
    for v in dv.values():
        card.add_data_validation(v)

    _write_guide(guide, labeler, len(rows))

    for col, width in {"A": 13, "B": 96, "C": 34, "D": 46, "E": 4}.items():
        card.column_dimensions[col].width = width
    card.column_dimensions["E"].hidden = True
    card.freeze_panes = "A2"
    card.append(["", "Mỗi khối là một tuyên bố. Điền các ô VÀNG ở cột C (chọn trong danh sách). "
                 "Ô trạng thái xanh ✓ là xong.", "Trả lời", "Gợi ý nhanh", "key"])
    for c in card[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = WRAP
    card.row_dimensions[1].height = 32

    status_cells: list[tuple[int, str, str]] = []
    answer_cells: list[str] = []
    r = 2
    for row in rows:
        idx = row["idx"]
        head = r
        where = (f"{row['ticker']} · {DOC_TYPES.get(row['document_type'], row['document_type'])} "
                 f"· năm {row.get('publication_year') or '?'} · trang {row.get('unit_index')} "
                 f"· nhóm {SPLITS.get(row.get('split'), row.get('split'))}")
        _put(card, r, ["", where, None, "", f"{idx}|_head"])
        card.cell(r, 1, f"CÂU {idx + 1}/{len(rows)}")
        for c in card[r]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = FILL_HEAD
        r += 1

        text = clean(row["text"])
        _put(card, r, ["Tuyên bố", text, "", "", f"{idx}|_claim"])
        card.cell(r, 2).font = Font(bold=True, size=11)
        for c in card[r][:4]:
            c.fill = FILL_CLAIM
        card.row_dimensions[r].height = _height(text, 100)
        r += 1

        q1 = r
        _answer(card, r, "Q1", "Câu này có phải tuyên bố môi trường của CHÍNH doanh nghiệp không?",
                "Số vĩ mô, kế hoạch kinh doanh, chỉ tiêu xã hội, tiêu đề mục, tin tức → no. "
                "Chọn no là xong câu này.", f"{idx}|is_claim", dv["is_claim"])
        answer_cells.append(f"C{r}")
        r += 1
        typ = r
        _answer(card, r, "Loại", "Tuyên bố thuộc loại nào?",
                "Chọn loại cụ thể nhất. 'chung chung' chỉ khi không có chủ đề nào khác.",
                f"{idx}|claim_type", dv["claim_type"])
        answer_cells.append(f"C{r}")
        r += 1
        first_attr = r
        for attr in ATTRS:
            label, hint = ATTR_QUESTIONS[attr]
            _answer(card, r, f"Q2 · {attr}", label, hint, f"{idx}|attrs.{attr}", dv["attr"])
            answer_cells.append(f"C{r}")
            r += 1
        last_attr = r - 1

        first_ev = r
        for ev in row["evidence_candidates"]:
            full = clean(ev["text"])
            shown = full if len(full) <= DISPLAY_LIMIT else full[:DISPLAY_LIMIT] + " …[cắt bớt]"
            source = (f"[{DOC_TYPES.get(ev['document_type'], ev['document_type'])} "
                      f"{ev.get('publication_year') or '?'} · trang {ev.get('unit_index')}]  ")
            _answer(card, r, f"E{ev['rank']}", source + shown,
                    "Đoạn này so với tuyên bố: xác nhận / một phần / mâu thuẫn / bối cảnh / "
                    "không liên quan?" if r == first_ev else "",
                    f"{idx}|evidence.{ev['evidence_id']}", dv["relation"])
            card.row_dimensions[r].height = _height(source + shown, 100)
            answer_cells.append(f"C{r}")
            r += 1
        last_ev = r - 1

        verdict = r
        _answer(card, r, "Q4", "Kết luận: bằng chứng ở trên đủ tới mức nào cho tuyên bố này?",
                "Nhãn là MỨC ĐỦ CỦA BẰNG CHỨNG, không phải \"DN có greenwashing\". "
                "Không có đoạn nào liên quan → INSUFFICIENT_EVIDENCE.",
                f"{idx}|verdict", dv["verdict"])
        answer_cells.append(f"C{r}")
        r += 1
        notes = r
        _put(card, r, ["Ghi chú", "", "", "Tuỳ chọn. Bắt buộc khi chọn 'không chắc'. "
                       "Mâu thuẫn nội bộ: bắt đầu bằng INCONSISTENCY:", f"{idx}|notes"])
        card.cell(r, 2).fill = PatternFill("solid", fgColor="FFFFFF")
        card.cell(r, 2).border = BOX
        card.row_dimensions[r].height = 30
        r += 1

        ev_rng = f"C{first_ev}:C{last_ev}"
        v = f"C{verdict}"
        formula = (
            f'=IF(C{q1}="","… chưa làm",'
            f'IF(LEFT(C{q1},2)="no","✓ Xong (không phải tuyên bố)",'
            f'IF(LEFT(C{q1},2)="un",IF(B{notes}="","⚠ Không chắc: ghi lý do ở ô Ghi chú",'
            f'"✓ Xong (không chắc)"),'
            f'IF(OR(C{typ}="",COUNTBLANK(C{first_attr}:C{last_attr})>0),"⚠ Thiếu loại hoặc thuộc tính",'
            f'IF(COUNTBLANK({ev_rng})>0,"⚠ Thiếu nhãn cho đoạn bằng chứng",'
            f'IF({v}="","⚠ Thiếu kết luận Q4",'
            f'IF(AND(LEFT({v},9)="SUPPORTED",COUNTIF({ev_rng},"supports*")=0),'
            f'"⚠ SUPPORTED cần ít nhất 1 đoạn supports",'
            f'IF(AND(LEFT({v},9)="SUPPORTED",COUNTIF({ev_rng},"contradicts*")>0),'
            f'"⚠ SUPPORTED nhưng còn đoạn contradicts",'
            f'IF(AND(LEFT({v},12)="CONTRADICTED",COUNTIF({ev_rng},"contradicts*")=0),'
            f'"⚠ CONTRADICTED cần ít nhất 1 đoạn contradicts",'
            f'IF(AND(LEFT({v},12)="INSUFFICIENT",COUNTIF({ev_rng},"supports*")'
            f'+COUNTIF({ev_rng},"partial*")+COUNTIF({ev_rng},"contradicts*")>0),'
            f'"⚠ INSUFFICIENT nhưng có đoạn liên quan — xem lại","✓ Xong"))))))))))'
        )
        card.cell(head, 3, formula)
        card.cell(head, 3).font = Font(bold=True, color="FFFFFF")
        status_cells.append((idx, row["ticker"], f"C{head}"))
        card.append([])
        r += 1

    for ref in answer_cells:
        card.conditional_formatting.add(ref, FormulaRule(formula=[f'ISBLANK({ref})'], fill=FILL_TODO))

    progress.column_dimensions["A"].width = 8
    progress.column_dimensions["B"].width = 10
    progress.column_dimensions["C"].width = 60
    progress.append(["Người gán", labeler])
    progress.append(["Đã xong", f'=COUNTIF(C4:C{3 + len(status_cells)},"✓*")&" / {len(status_cells)}"'])
    progress.append(["Câu", "DN", "Trạng thái (bấm vào tên câu ở sheet Gán nhãn để sửa)"])
    for c in progress[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
    for idx, ticker, ref in status_cells:
        progress.append([idx + 1, ticker, f"='Gán nhãn'!{ref}"])
    last = 3 + len(status_cells)
    progress.conditional_formatting.add(
        f"C4:C{last}", FormulaRule(formula=['LEFT(C4,1)="✓"'], fill=FILL_OK))
    progress.conditional_formatting.add(
        f"C4:C{last}", FormulaRule(formula=['LEFT(C4,1)="⚠"'], fill=FILL_WARN))
    progress["B1"].font = Font(bold=True)
    progress["B2"].font = Font(bold=True, color=GREEN, size=13)

    wb.active = 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def dropdowns(wb, spec: dict[str, tuple[str, list[str]]], sheet: str = "DS") -> dict:
    """List validations backed by named ranges on a hidden sheet.

    Excel before 2010 rejects a validation list that points straight at another
    sheet; a defined name works in every version and in Google Sheets.
    """
    out = {}
    for key, (col, values) in spec.items():
        name = f"L_{key}"
        wb.defined_names[name] = DefinedName(name, attr_text=f"{sheet}!${col}$1:${col}${len(values)}")
        out[key] = DataValidation(type="list", formula1=f"={name}", allow_blank=True,
                                  showErrorMessage=True, errorTitle="Chọn trong danh sách",
                                  error="Hãy chọn một giá trị trong danh sách thả xuống.")
    return out


def _put(ws, r: int, values: list) -> None:
    for c, value in enumerate(values, start=1):
        cell = ws.cell(r, c, value)
        cell.alignment = WRAP


def _answer(ws, r, tag, question, hint, key, validation) -> None:
    _put(ws, r, [tag, question, None, hint, key])
    ws.cell(r, 1).font = Font(bold=True, color=GREEN)
    ws.cell(r, 3).border = BOX
    ws.cell(r, 4).font = Font(size=9, color="5B6770")
    validation.add(f"C{r}")


def _write_guide(ws, labeler: str, n: int) -> None:
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Bạn là", labeler),
        ("Số tuyên bố", str(n)),
        ("Cách làm", "Sang sheet 'Gán nhãn'. Mỗi khối là một tuyên bố: đọc TUYÊN BỐ, trả lời Q1. "
                     "Nếu Q1 = no thì xong, sang khối sau. Nếu yes: chọn Loại, 5 thuộc tính Q2, "
                     "nhãn cho từng đoạn E1…E5, rồi kết luận Q4."),
        ("Ô vàng", "là ô chưa điền. Ô trạng thái ở đầu mỗi khối tự chuyển ✓ khi đủ và hợp lệ."),
        ("Sheet Tiến độ", "cho biết còn câu nào chưa xong hoặc có cảnh báo ⚠."),
        ("Làm độc lập", "Nếu lô này có hai người gán: KHÔNG bàn với người kia trước khi cả hai nộp. "
                        "Độ đồng thuận (κ) chỉ có nghĩa khi hai người làm riêng."),
        ("", ""),
        ("3 lỗi hay gặp nhất", ""),
        ("1", "Số vĩ mô, kế hoạch kinh doanh, chỉ tiêu xã hội, tiêu đề mục → Q1 = no. "
              "Không phải tuyên bố môi trường của DN."),
        ("2", "Nhãn là MỨC ĐỦ CỦA BẰNG CHỨNG cho tuyên bố này, không phải 'DN có greenwashing'. "
              "Thiếu dữ liệu không làm DN thành xấu."),
        ("3", "INSUFFICIENT_EVIDENCE là câu trả lời đúng khi không đoạn nào xác nhận/bác bỏ, "
              "không phải né tránh."),
        ("", ""),
        ("Quan hệ bằng chứng", ""),
        ("supports", "xác nhận đúng chỉ số, cùng kỳ / phạm vi"),
        ("partial", "cùng chủ đề, xác nhận một phần (khác kỳ, chỉ số liên quan, chỉ nói có chương trình)"),
        ("contradicts", "số hoặc xu hướng khác trên CÙNG chỉ số CÙNG kỳ"),
        ("context", "liên quan nhưng không xác nhận/bác bỏ — kể cả khi đoạn đó lặp lại chính tuyên bố"),
        ("not_relevant", "không liên quan"),
        ("", ""),
        ("Kết luận Q4", ""),
        ("SUPPORTED", "≥ 1 đoạn supports (không phải chính câu tuyên bố), không còn contradicts"),
        ("PARTIALLY_SUPPORTED", "chỉ có partial; hoặc có supports nhưng thiếu ≥ 2 thuộc tính"),
        ("CONTRADICTED", "≥ 1 đoạn contradicts CÙNG chỉ số. Khác chỉ số thì không đủ — ghi INCONSISTENCY:"),
        ("UNSUPPORTED", "nguồn lẽ ra phải có số liệu (vd. mục 6 TT 96) nhưng không có"),
        ("INSUFFICIENT_EVIDENCE", "không đoạn nào supports / partial / contradicts"),
        ("", ""),
        ("Xong thì", "Lưu file (giữ nguyên tên) và gửi lại cho Nhân. Không đổi tên sheet, "
                     "không xoá hay chèn dòng."),
    ]
    for a, b in lines:
        ws.append([a, b])
    for row in ws.iter_rows():
        row[0].font = Font(bold=True, color=GREEN)
        row[1].alignment = WRAP
    for r in (8, 13, 20):
        ws.cell(r, 1).font = Font(bold=True, size=12, color="FFFFFF")
        ws.cell(r, 1).fill = FILL_HEAD


# ------------------------------------------------------------------- import


ALLOWED = {
    "is_claim": {"yes", "no", "unsure"},
    "claim_type": {code(v) for v in CLAIM_TYPES},
    "attr": {code(v) for v in ATTR_VALUES},
    "relation": {code(v) for v in RELATIONS},
    "verdict": {code(v) for v in VERDICTS},
}


def import_workbook(xlsx: Path, final: bool = False, out: Path | None = None) -> tuple[Path, list[str]]:
    wb = load_workbook(xlsx)
    meta = {row[0].value: row[1].value for row in wb["_meta"].iter_rows()}
    session = SESSIONS / meta["session"]
    labeler = meta["labeler"]
    rows = {r["idx"]: r for r in read_jsonl(session)}
    answers: dict[int, dict] = {i: {} for i in rows}

    ws = wb["Gán nhãn"]
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, 5).value
        if not key or "|" not in str(key):
            continue
        idx_text, field = str(key).split("|", 1)
        if field.startswith("_"):
            continue
        value = ws.cell(r, 2 if field == "notes" else 3).value
        answers[int(idx_text)][field] = value

    problems: list[str] = []
    stamp = dt.datetime.fromtimestamp(xlsx.stat().st_mtime, dt.UTC).isoformat(timespec="seconds")
    for idx, row in rows.items():
        got = answers[idx]
        labels = {"is_claim": None, "claim_type": None, "attrs": {a: None for a in ATTRS},
                  "evidence": {}, "verdict": None, "notes": (got.get("notes") or "").strip(),
                  "labeler": None, "labeled_at": None}
        where = f"câu {idx + 1} ({row['ticker']})"
        is_claim = code(got.get("is_claim"))
        if is_claim is None:
            problems.append(f"{where}: chưa trả lời Q1")
            row["labels"] = labels
            continue
        if not _check(is_claim, "is_claim", where, "Q1", problems):
            row["labels"] = labels
            continue
        labels["is_claim"] = is_claim
        labels["labeler"] = labeler
        labels["labeled_at"] = stamp
        if is_claim == "unsure" and not labels["notes"]:
            problems.append(f"{where}: chọn 'không chắc' nhưng không ghi lý do")
        if is_claim == "yes":
            ctype = code(got.get("claim_type"))
            if _check(ctype, "claim_type", where, "Loại", problems):
                labels["claim_type"] = ctype
            for attr in ATTRS:
                value = code(got.get(f"attrs.{attr}"))
                if _check(value, "attr", where, f"thuộc tính {attr}", problems):
                    labels["attrs"][attr] = value
            for cand in row["evidence_candidates"]:
                value = code(got.get(f"evidence.{cand['evidence_id']}"))
                if _check(value, "relation", where, f"E{cand['rank']}", problems):
                    labels["evidence"][cand["evidence_id"]] = value
            verdict = code(got.get("verdict"))
            if _check(verdict, "verdict", where, "Q4", problems):
                labels["verdict"] = verdict
                problems.extend(f"{where}: {p}" for p in consistency(verdict, labels["evidence"].values()))
        row["labels"] = labels

    ordered = [rows[i] for i in sorted(rows)]
    if out is None:
        out = session if final else LABELER_COPIES / f"{session.stem}_{labeler}.jsonl"
    write_jsonl(out, ordered)
    return out, problems


def _check(value, kind: str, where: str, what: str, problems: list[str]) -> bool:
    if value is None:
        problems.append(f"{where}: {what} còn trống")
        return False
    if value not in ALLOWED[kind]:
        problems.append(f"{where}: {what} = '{value}' không nằm trong danh sách")
        return False
    return True


def consistency(verdict: str, relations) -> list[str]:
    """The same rules the workbook shows live, so a file edited outside Excel is held to them too."""
    counts = Counter(relations)
    related = counts["supports"] + counts["partial"] + counts["contradicts"]
    issues = []
    if verdict == "SUPPORTED" and not counts["supports"]:
        issues.append("SUPPORTED nhưng không đoạn nào supports")
    if verdict == "SUPPORTED" and counts["contradicts"]:
        issues.append("SUPPORTED nhưng còn đoạn contradicts")
    if verdict == "CONTRADICTED" and not counts["contradicts"]:
        issues.append("CONTRADICTED nhưng không đoạn nào contradicts")
    if verdict == "INSUFFICIENT_EVIDENCE" and related:
        issues.append("INSUFFICIENT_EVIDENCE nhưng có đoạn supports/partial/contradicts")
    return issues


# -------------------------------------------------------------------- kappa


def cohen_kappa(pairs: list[tuple[str, str]]) -> dict:
    n = len(pairs)
    if not n:
        return {"n": 0, "kappa": None, "agreement": None}
    agree = sum(a == b for a, b in pairs) / n
    left, right = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    expected = sum(left[c] * right[c] for c in set(left) | set(right)) / (n * n)
    if expected >= 1.0:
        # One category on both sides: agreement is trivially total and kappa is undefined.
        return {"n": n, "kappa": None, "agreement": round(agree, 4),
                "note": "chỉ một giá trị ở cả hai phía — κ không xác định"}
    return {"n": n, "kappa": round((agree - expected) / (1 - expected), 4),
            "agreement": round(agree, 4)}


def kappa(a_path: Path, b_path: Path) -> tuple[dict, list[dict]]:
    a_rows = {r["claim_id"]: r for r in read_jsonl(a_path)}
    b_rows = {r["claim_id"]: r for r in read_jsonl(b_path)}
    shared = [cid for cid in a_rows if cid in b_rows]
    fields: dict[str, list[tuple[str, str]]] = {"is_claim": [], "claim_type": [], "verdict": []}
    fields.update({f"attrs.{a}": [] for a in ATTRS})
    fields["evidence"] = []
    disagreements: list[dict] = []

    def note(row, field, x, y):
        if x != y:
            disagreements.append({"idx": row["idx"] + 1, "ticker": row["ticker"], "field": field,
                                  "a": x, "b": y, "text": clean(row["text"])[:300]})

    for cid in shared:
        ra, rb = a_rows[cid], b_rows[cid]
        la, lb = ra["labels"], rb["labels"]
        if not (la.get("is_claim") and lb.get("is_claim")):
            continue
        fields["is_claim"].append((la["is_claim"], lb["is_claim"]))
        note(ra, "Q1 is_claim", la["is_claim"], lb["is_claim"])
        if la["is_claim"] != "yes" or lb["is_claim"] != "yes":
            continue
        for key in ("claim_type", "verdict"):
            if la.get(key) and lb.get(key):
                fields[key].append((la[key], lb[key]))
                note(ra, key, la[key], lb[key])
        for attr in ATTRS:
            x, y = la["attrs"].get(attr), lb["attrs"].get(attr)
            if x and y:
                fields[f"attrs.{attr}"].append((x, y))
                note(ra, f"attrs.{attr}", x, y)
        for cand in ra["evidence_candidates"]:
            eid = cand["evidence_id"]
            x, y = la["evidence"].get(eid), lb["evidence"].get(eid)
            if x and y:
                fields["evidence"].append((x, y))
                note(ra, f"E{cand['rank']}", x, y)

    names = (a_rows[shared[0]]["labels"].get("labeler") if shared else "A",
             b_rows[shared[0]]["labels"].get("labeler") if shared else "B")
    return {"a": names[0] or "A", "b": names[1] or "B", "claims": len(shared),
            "fields": {k: cohen_kappa(v) for k, v in fields.items()}}, disagreements


def write_kappa_report(result: dict, disagreements: list[dict], stem: str) -> tuple[Path, Path]:
    BENCHMARK.mkdir(exist_ok=True)
    a, b = result["a"], result["b"]
    lines = [
        f"# Đồng thuận giữa {a} và {b} — {stem}", "",
        f"Tuyên bố chung: {result['claims']} · ngưỡng cam kết Vòng 1: **κ ≥ 0,7**", "",
        "| Trường | n | Đồng ý | κ |", "| --- | ---: | ---: | ---: |",
    ]
    for field, m in result["fields"].items():
        k = "—" if m["kappa"] is None else f"{m['kappa']:.3f}"
        agree = "—" if m["agreement"] is None else f"{m['agreement'] * 100:.0f}%"
        lines.append(f"| {field} | {m['n']} | {agree} | {k} |")
    lines += ["", f"## {len(disagreements)} chỗ khác nhau — mang vào buổi trọng tài", ""]
    for d in disagreements:
        lines.append(f"- câu {d['idx']} ({d['ticker']}) · `{d['field']}`: {a} = `{d['a']}` · {b} = `{d['b']}`")
    md = BENCHMARK / f"kappa_{stem}.md"
    md.write_text("\n".join(lines) + "\n", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Khác nhau"
    ws.append(["Câu", "DN", "Trường", a, b, "Thống nhất (điền trong buổi trọng tài)", "Tuyên bố"])
    for d in disagreements:
        ws.append([d["idx"], d["ticker"], d["field"], d["a"], d["b"], "", d["text"]])
    for col, width in zip("ABCDEFG", (6, 8, 18, 22, 22, 30, 90), strict=True):
        ws.column_dimensions[col].width = width
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
    for row in ws.iter_rows(min_row=2):
        row[6].alignment = WRAP
    xlsx = BENCHMARK / f"kappa_{stem}_khac_nhau.xlsx"
    wb.save(xlsx)
    return md, xlsx


# ------------------------------------------------ the other returned sheets


def _header_row(ws, first_cell: str) -> int:
    for r in range(1, 15):
        if str(ws.cell(r, 1).value or "").strip() == first_cell:
            return r
    raise ValueError(f"{ws.title}: header row starting with {first_cell!r} not found")


def _columns(ws, row: int) -> dict[str, int]:
    return {str(ws.cell(row, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}


def import_hpg_queue(xlsx: Path) -> Path:
    """Quỳnh's yes/no on each item of the HPG review queue -> benchmark/hpg_queue_review.json."""
    ws = load_workbook(xlsx)["Hàng đợi HPG"]
    head = _header_row(ws, "#")
    cols = _columns(ws, head)
    answer_col = next(c for name, c in cols.items() if name.startswith("Có phải"))
    why_col = next(c for name, c in cols.items() if name.startswith("Nếu Không"))
    note_col = next(c for name, c in cols.items() if name.startswith("Một dòng"))
    items = []
    for r in range(head + 1, ws.max_row + 1):
        rank = ws.cell(r, 1).value
        if not isinstance(rank, int):
            continue
        answer = (ws.cell(r, answer_col).value or "").strip() or None
        items.append({
            "rank": rank,
            "required": ws.cell(r, cols["Bắt buộc"]).value == "bắt buộc",
            "kind": ws.cell(r, cols["Loại mục"]).value,
            "page": ws.cell(r, cols["Trang"]).value,
            "text": ws.cell(r, cols["Nội dung hệ thống đưa lên hàng đợi"]).value,
            "item_id": ws.cell(r, cols["item_id"]).value,
            "is_claim": {"Có": "yes", "Không": "no", "Không chắc": "unsure"}.get(answer),
            "why_not": ws.cell(r, why_col).value,
            "note": ws.cell(r, note_col).value,
        })

    def share(subset):
        answered = [i for i in subset if i["is_claim"] in {"yes", "no"}]
        return {"answered": len(answered), "yes": sum(i["is_claim"] == "yes" for i in answered),
                "precision": round(sum(i["is_claim"] == "yes" for i in answered) / len(answered), 4)
                if answered else None}

    report = {
        "source": xlsx.name,
        "run_id": "185480a0c8b747c4",
        "imported_at": dt.datetime.now().isoformat(timespec="seconds"),
        "top14": share([i for i in items if i["rank"] <= 14]),
        "all": share(items),
        "why_not": dict(Counter(i["why_not"] for i in items if i["is_claim"] == "no" and i["why_not"])),
        "items": items,
    }
    out = BENCHMARK / "hpg_queue_review.json"
    BENCHMARK.mkdir(exist_ok=True)
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def _minutes(start, end) -> float | None:
    def as_minutes(value):
        if value is None or value == "":
            return None
        if isinstance(value, dt.time):
            return value.hour * 60 + value.minute + value.second / 60
        if isinstance(value, dt.datetime):
            return value.hour * 60 + value.minute + value.second / 60
        if isinstance(value, int | float):
            return float(value) * 24 * 60          # Excel stores a time as a fraction of a day
        text = str(value).strip().replace("h", ":")
        hours, _, minutes = text.partition(":")
        return int(hours) * 60 + int(minutes or 0)
    a, b = as_minutes(start), as_minutes(end)
    if a is None or b is None:
        return None
    return round(b - a if b >= a else b + 24 * 60 - a, 1)


def import_timing(xlsx: Path) -> Path:
    """The timing study (manual vs GreenScan) -> benchmark/timing_study.json."""
    ws = load_workbook(xlsx)["Bài đo"]
    head = _header_row(ws, "#")
    cols = _columns(ws, head)

    def col(prefix):
        return next(c for name, c in cols.items() if name.startswith(prefix))

    items = []
    for r in range(head + 1, ws.max_row + 1):
        number = ws.cell(r, 1).value
        if not isinstance(number, int):
            continue
        minutes = _minutes(ws.cell(r, col("Bắt đầu")).value, ws.cell(r, col("Kết thúc")).value)
        correct = ws.cell(r, col("Số đoạn ĐÚNG")).value
        items.append({
            "n": number,
            "group": "manual" if str(ws.cell(r, col("Nhóm")).value).startswith("A") else "greenscan",
            "claim": ws.cell(r, col("Tuyên bố")).value,
            "minutes": minutes,
            "pages_found": ws.cell(r, col("Trang bằng chứng")).value,
            "passages_found": ws.cell(r, col("Số đoạn bằng chứng")).value,
            "passages_correct": correct if isinstance(correct, int | float) else None,
            "verdict": ws.cell(r, col("Kết luận")).value,
            "confidence": ws.cell(r, col("Tự tin")).value,
            "note": ws.cell(r, col("Ghi chú")).value,
        })
    done = [i for i in items if i["minutes"] is not None]
    report = {
        "source": xlsx.name,
        "imported_at": dt.datetime.now().isoformat(timespec="seconds"),
        "completed": len(done),
        "items": items,
    }
    out = BENCHMARK / "timing_study.json"
    BENCHMARK.mkdir(exist_ok=True)
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def import_all(folder: Path) -> list[str]:
    """Import every workbook Quỳnh and Thảo sent back, whatever order they arrive in.

    Gold lô 2 and lô 3 have one labeller, so their import is final. Lô 1 is
    labelled twice: both copies are kept apart and κ is computed as soon as both
    are in; it becomes gold only after adjudication (`import --final` on the
    reconciled workbook).
    """
    log = []
    for xlsx in sorted(folder.glob("*.xlsx")):
        if xlsx.name.startswith("~$"):
            continue
        sheets = load_workbook(xlsx, read_only=True).sheetnames
        if "_meta" in sheets:
            meta = {r[0].value: r[1].value for r in load_workbook(xlsx)["_meta"].iter_rows()}
            copy, problems = import_workbook(xlsx)
            log.append(f"{xlsx.name}: {meta['labeler']} -> {copy.relative_to(ROOT)}, {len(problems)} vấn đề")
            log.extend(f"    - {p}" for p in problems[:30])
            if "kappa" not in meta["session"]:
                final, _ = import_workbook(xlsx, final=True)
                log.append(f"    gold (một người gán): {final.relative_to(ROOT)}")
        elif "Hàng đợi HPG" in sheets:
            log.append(f"{xlsx.name}: hàng đợi HPG -> {import_hpg_queue(xlsx).relative_to(ROOT)}")
        elif "Bài đo" in sheets:
            log.append(f"{xlsx.name}: bài đo thời gian -> {import_timing(xlsx).relative_to(ROOT)}")
        else:
            log.append(f"{xlsx.name}: không nhận ra loại file, bỏ qua")
    for a in sorted(LABELER_COPIES.glob("*_quynh.jsonl")):
        b = a.with_name(a.name.replace("_quynh.jsonl", "_thao.jsonl"))
        if b.exists():
            result, disagreements = kappa(a, b)
            md, _ = write_kappa_report(result, disagreements, a.stem.rsplit("_", 1)[0])
            log.append(f"κ {a.stem.rsplit('_', 1)[0]}: {md.relative_to(ROOT)} · "
                       f"{len(disagreements)} chỗ khác nhau cho buổi trọng tài")
    return log


# --------------------------------------------------------------------- main


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("export")
    s.add_argument("session")
    s.add_argument("--labeler", required=True)
    s.add_argument("--out", required=True)
    s = sub.add_parser("import")
    s.add_argument("xlsx")
    s.add_argument("--final", action="store_true",
                   help="write the adjudicated gold where evaluate_gold.py reads it")
    s = sub.add_parser("kappa")
    s.add_argument("a")
    s.add_argument("b")
    s = sub.add_parser("import-all", help="every returned workbook in a folder")
    s.add_argument("folder")
    args = ap.parse_args()

    if args.cmd == "import-all":
        for line in import_all(Path(args.folder)):
            print(line)
        return

    if args.cmd == "export":
        path = export(Path(args.session), args.labeler, Path(args.out))
        print(f"written: {path}")
    elif args.cmd == "import":
        path, problems = import_workbook(Path(args.xlsx), final=args.final)
        print(f"written: {path}")
        if problems:
            print(f"{len(problems)} vấn đề cần người gán sửa:")
            for p in problems:
                print("  -", p)
        else:
            print("không có vấn đề nào")
    elif args.cmd == "kappa":
        result, disagreements = kappa(Path(args.a), Path(args.b))
        stem = Path(args.a).stem.rsplit("_", 1)[0]
        md, xlsx = write_kappa_report(result, disagreements, stem)
        print(md.read_text(encoding="utf-8"))
        print(f"written: {md} · {xlsx}")


if __name__ == "__main__":
    main()
