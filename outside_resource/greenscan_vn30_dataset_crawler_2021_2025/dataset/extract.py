from pathlib import Path
import fitz
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook

def extract(path:Path):
    s=path.suffix.lower()
    if s=='.pdf':
        doc=fitz.open(path); return [{'unit_type':'page','unit_index':i+1,'text':p.get_text('text')} for i,p in enumerate(doc)]
    if s=='.html':
        soup=BeautifulSoup(path.read_text(encoding='utf-8',errors='replace'),'lxml')
        for tag in soup(['script','style','noscript','svg']): tag.decompose()
        text='\n'.join(x.get_text(' ',strip=True) for x in soup.select('title,h1,h2,h3,h4,p,li,td,th') if x.get_text(' ',strip=True)); return [{'unit_type':'html','unit_index':1,'text':text}]
    if s=='.docx':
        d=Document(path); text='\n'.join(p.text for p in d.paragraphs if p.text.strip()); return [{'unit_type':'document','unit_index':1,'text':text}]
    if s=='.xlsx':
        wb=load_workbook(path,read_only=True,data_only=True); out=[]
        for i,sh in enumerate(wb.worksheets,1):
            rows=['\t'.join(str(v) for v in row if v not in (None,'')) for row in sh.iter_rows(values_only=True)]; out.append({'unit_type':'sheet','unit_index':i,'title':sh.title,'text':'\n'.join(x for x in rows if x)})
        return out
    if s in {'.csv','.txt'}: return [{'unit_type':'document','unit_index':1,'text':path.read_text(encoding='utf-8',errors='replace')}]
    return []
